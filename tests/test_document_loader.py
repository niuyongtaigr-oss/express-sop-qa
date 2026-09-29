"""文档解析层测试 (无需 Ollama / 网络)

覆盖:
  - 中文编码识别 (utf-8 / GBK / BOM) —— 国内资料按 utf-8 硬读会静默出乱码
  - 文本清洗 (控制字符 / 汉字间隔 / 空白压缩)
  - 表格 → 带表头上下文的行文本
  - 各格式解析: txt / md / csv / docx / xlsx / pdf
  - 体积、字符数兜底与截断标记
  - doc_id 派生 (ASCII 直用 / 非 ASCII 哈希, 避免中文名撞 id)
  - DocumentParseError → HTTP 400 的映射

样本文件均在测试内构造, 不依赖仓库内的二进制 fixture。
"""

import io

import pytest

from app.core.exceptions import AppError, DocumentParseError
from app.infrastructure.document_loader import (
    CsvParser,
    DocumentLoader,
    ExtractedDocument,
    PlainTextParser,
    decode_text_bytes,
    normalize_text,
    render_rows_as_text,
)
from app.services.rag_service import RagService, derive_doc_id


# ── 编码识别 ─────────────────────────────────────────────


def test_decode_utf8():
    assert decode_text_bytes("破损件处理".encode("utf-8")) == ("破损件处理", "utf-8")


def test_decode_gbk_becomes_gb18030():
    """GBK 资料必须能读出来 —— 按 utf-8 硬读会抛错或被静默替换成乱码"""
    text, encoding = decode_text_bytes("破损件处理".encode("gbk"))
    assert text == "破损件处理"
    assert encoding == "gb18030"


def test_decode_respects_utf8_bom():
    assert decode_text_bytes("破损".encode("utf-8-sig")) == ("破损", "utf-8-sig")


def test_decode_respects_utf16_bom():
    assert decode_text_bytes("破损".encode("utf-16")) == ("破损", "utf-16")


# ── 文本清洗 ─────────────────────────────────────────────


def test_normalize_strips_control_chars_and_exotic_spaces():
    # \x00 是 pypdf 对某些字体/编码的常见产物; \u00a0/\u3000 是排版空格
    assert normalize_text("a\x00b\u00a0c\u3000d") == "ab c d"


def test_normalize_merges_han_gaps():
    """汉字被字距误判拆开时必须并回, 否则检索切词全废"""
    assert normalize_text("理 赔 流 程") == "理赔流程"
    assert normalize_text("破 损 件 处 理 规 范") == "破损件处理规范"
    # 只并汉字之间的空格; 拉丁字母侧的空格属真实语义 (见下个用例)
    assert normalize_text("SOP 文 档 说明") == "SOP 文档说明"


def test_normalize_keeps_latin_spacing():
    # 拉丁字母与数字之间的空格是真实语义, 不能并掉
    assert normalize_text("Java 21 / Spring Boot") == "Java 21 / Spring Boot"


def test_normalize_unifies_newlines_and_blank_lines():
    assert normalize_text("a\r\n\r\n\r\n\r\nb") == "a\n\nb"


# ── 表格渲染 ─────────────────────────────────────────────


def test_render_rows_prefixes_header():
    rows = [["网点", "破损数"], ["上海浦东", "3"]]
    assert render_rows_as_text(rows) == "网点: 上海浦东; 破损数: 3"


def test_render_rows_header_only():
    assert render_rows_as_text([["A", "B"]]) == "A | B"


def test_render_rows_skips_blank_rows():
    assert render_rows_as_text([["A"], ["", ""], ["v"]]) == "A: v"


def test_render_rows_empty():
    assert render_rows_as_text([]) == ""
    assert render_rows_as_text([["", ""]]) == ""


# ── 路由与兜底 ───────────────────────────────────────────


def test_supported_extensions_cover_common_office_formats():
    exts = DocumentLoader().supported_extensions()
    for expected in (".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md"):
        assert expected in exts


def test_rejects_unsupported_extension():
    with pytest.raises(DocumentParseError, match="不支持的文件类型"):
        DocumentLoader().load("payload.exe", b"MZ")


def test_rejects_empty_file():
    with pytest.raises(DocumentParseError, match="内容为空"):
        DocumentLoader().load("a.txt", b"")


def test_rejects_oversize_file():
    with pytest.raises(DocumentParseError, match="文件过大"):
        DocumentLoader(max_bytes=10).load("a.txt", b"x" * 20)


def test_truncates_and_flags():
    doc = DocumentLoader(max_chars=5).load("a.txt", "一二三四五六七八".encode("utf-8"))
    assert doc.text == "一二三四五"
    assert doc.metadata["truncated"] is True
    assert doc.metadata["chars_raw"] == 8
    assert doc.metadata["chars"] == 5  # 截断后按字符计


def test_attaches_source_metadata():
    doc = DocumentLoader().load("台账.csv", "A,B\n1,2\n".encode("utf-8"))
    assert doc.metadata["filename"] == "台账.csv"
    assert doc.metadata["ext"] == ".csv"
    assert doc.metadata["bytes"] > 0
    assert doc.metadata["truncated"] is False


# ── 各格式解析 ───────────────────────────────────────────


def test_txt_parser_gbk():
    doc = PlainTextParser().parse("x.txt", "破损件需备案".encode("gbk"))
    assert "破损件需备案" in doc.text
    assert doc.metadata["encoding"] == "gb18030"


def test_txt_parser_rejects_blank():
    with pytest.raises(DocumentParseError):
        PlainTextParser().parse("x.txt", b"   \n\n  ")


def test_csv_parser_header_prefixed_rows():
    doc = CsvParser().parse("x.csv", "网点,件数\n上海浦东,3\n".encode("utf-8"))
    assert doc.text == "网点: 上海浦东; 件数: 3"
    assert doc.metadata["rows"] == 2


def test_csv_parser_rejects_empty():
    with pytest.raises(DocumentParseError):
        CsvParser().parse("x.csv", b"\n\n")


def _docx_bytes() -> bytes:
    """测试内构造 .docx (含段落 + 表格), 不依赖二进制 fixture"""
    import docx  # python-docx

    document = docx.Document()
    document.add_paragraph("理赔流程说明")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "环节", "时限"
    table.cell(1, 0).text, table.cell(1, 1).text = "异常登记", "立即"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_docx_parser_paragraphs_and_tables():
    doc = DocumentLoader().load("制度.docx", _docx_bytes())
    assert "理赔流程说明" in doc.text
    assert "环节: 异常登记; 时限: 立即" in doc.text
    assert doc.metadata["tables"] == 1


def test_docx_parser_rejects_non_docx():
    with pytest.raises(DocumentParseError, match="Word 解析失败"):
        DocumentLoader().load("old.docx", b"not a docx at all")


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "赔付标准"
    sheet.append(["保价情况", "赔偿上限"])
    sheet.append(["未保价", "运费3倍"])
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def test_xlsx_parser_includes_sheet_name():
    doc = DocumentLoader().load("赔付.xlsx", _xlsx_bytes())
    assert "【工作表: 赔付标准】" in doc.text
    assert "保价情况: 未保价; 赔偿上限: 运费3倍" in doc.text
    assert doc.metadata["sheets"] == 1


def _pdf_bytes(text: str = "SOP rule text") -> bytes:
    """构造带文本层的最小 PDF (正确 xref), 用于 PDF 解析正常路径"""
    content = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode("latin-1")
    bodies = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>stream\n%s\nendstream" % (len(content), content),
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(bodies, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj" % i + body + b"endobj\n"
    xref_pos = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(bodies) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer<</Root 1 0 R/Size %d>>\nstartxref\n%d\n%%%%EOF\n" % (
        len(bodies) + 1,
        xref_pos,
    )
    return bytes(out)


def test_pdf_parser_extracts_text():
    doc = DocumentLoader().load("a.pdf", _pdf_bytes())
    assert "SOP rule text" in doc.text
    assert doc.metadata["pages"] == 1
    assert doc.metadata["pages_with_text"] == 1


def test_pdf_parser_rejects_scanned_pdf():
    """扫描件 (无文本层) 必须明确报错, 不能静默产出空文档"""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(DocumentParseError, match="扫描件"):
        DocumentLoader().load("scan.pdf", buf.getvalue())


def test_pdf_parser_rejects_broken_file():
    with pytest.raises(DocumentParseError, match="PDF 解析失败"):
        DocumentLoader().load("broken.pdf", b"%PDF-1.4 not really a pdf")


# ── doc_id 派生 ──────────────────────────────────────────


@pytest.mark.parametrize("name,expected", [("sop.txt", "sop"), ("my-doc_v2.pdf", "my-doc_v2")])
def test_derive_doc_id_keeps_ascii_stem(name, expected):
    assert derive_doc_id(name) == expected


def test_derive_doc_id_hashes_non_ascii():
    """中文文件名若不哈希, 会被向量库压成同样的下划线串而互相顶掉"""
    a, b = derive_doc_id("制度.docx"), derive_doc_id("赔付标准.xlsx")
    assert a.startswith("doc_") and b.startswith("doc_")
    assert a != b


def test_derive_doc_id_is_stable():
    assert derive_doc_id("制度.docx") == derive_doc_id("制度.docx")


# ── 异常 → HTTP 映射 ─────────────────────────────────────


def test_parse_error_is_app_error_400():
    """必须能被全局异常处理器接住转 400, 否则客户端拿到的是 500"""
    err = DocumentParseError("不支持的文件类型")
    assert isinstance(err, AppError)
    assert err.status_code == 400
    assert err.code == "document_parse_failed"


# ── RagService.ingest_upload 装配 ────────────────────────


class _RecordingStore:
    """只记录 add_document 调用的最小 VectorStore 替身"""

    def __init__(self):
        self.added: list[tuple] = []

    def add_document(self, doc_id, title, content, tenant_id="default"):
        self.added.append((doc_id, title, content, tenant_id))
        return 3


class _StubLoader:
    def __init__(self, doc: ExtractedDocument):
        self._doc = doc

    def load(self, filename, data):
        return self._doc

    def supported_extensions(self):
        return [".txt"]


def _settings():
    from app.config import Settings

    return Settings(_env_file=None)


def test_ingest_upload_uses_parsed_text_and_derives_id():
    store = _RecordingStore()
    parsed = ExtractedDocument(text="解析并清洗后的正文", metadata={"chars": 8})
    rag = RagService(store, None, _settings(), document_loader=_StubLoader(parsed))

    out = rag.ingest_upload("制度.docx", b"\x00\x01binary")

    assert out["doc_id"].startswith("doc_")     # 中文名 → 哈希 id
    assert out["title"] == "制度"
    assert out["chunks"] == 3
    assert store.added == [(out["doc_id"], "制度", "解析并清洗后的正文", "default")]


def test_ingest_upload_propagates_truncated_flag():
    """truncated 必须来自解析结果

    不能只断言 `out["truncated"] is False` —— 那是"两种实现都成立"的默认值断言:
    解析器没给 truncated 时它自然也是 False, 于是把透传写成硬编码 False 测试照样绿,
    而"文件被截断"这个**该报警**的信号就再也传不出来了。
    """
    store = _RecordingStore()
    parsed = ExtractedDocument(text="被截断的正文", metadata={"truncated": True})
    rag = RagService(store, None, _settings(), document_loader=_StubLoader(parsed))

    assert rag.ingest_upload("大文件.txt", b"x")["truncated"] is True

    store2 = _RecordingStore()
    ok = ExtractedDocument(text="完整正文", metadata={"truncated": False})
    rag2 = RagService(store2, None, _settings(), document_loader=_StubLoader(ok))
    assert rag2.ingest_upload("小文件.txt", b"x")["truncated"] is False

    # 解析器完全没提 truncated 时按"未截断"处理
    store3 = _RecordingStore()
    silent = ExtractedDocument(text="正文", metadata={})
    rag3 = RagService(store3, None, _settings(), document_loader=_StubLoader(silent))
    assert rag3.ingest_upload("a.txt", b"x")["truncated"] is False


def test_ingest_upload_respects_explicit_doc_id_and_tenant():
    store = _RecordingStore()
    rag = RagService(
        store, None, _settings(),
        document_loader=_StubLoader(ExtractedDocument(text="正文", metadata={})),
    )
    out = rag.ingest_upload("a.txt", b"x", doc_id="custom-1", title="自定义", tenant_id="t9")
    assert out["doc_id"] == "custom-1"
    assert store.added[0][3] == "t9"

"""文档解析层 — 把客户的实际资料 (PDF / Word / Excel / 文本) 转成可入库的纯文本

为什么需要这一层: RagService.add_document() 只接收纯文本, 而企业知识库的原始
资料是 PDF 规范、Word 制度、Excel 台账。没有解析层, 客户的资料就进不来 ——
这是「能演示」与「能交付」之间最大的一道坎。

设计要点:
  - 按扩展名路由。各解析器的第三方依赖全部懒加载, 缺失时给出可执行的报错
    (而非裸 ImportError 栈), 基础安装不必装齐全部解析库。
  - 中文编码自动识别 (BOM → utf-8 → gb18030): 国内资料 GBK 极常见, 直接按
    utf-8 读会产生乱码且不报错, 是最隐蔽的一类脏数据。
  - 表格 (xlsx/csv/docx 表格) 渲染成「表头: 值」的行文本, 比 markdown 表更
    适合 RAG 检索 —— 每行自带列名上下文, 切块后不会丢失语义。
  - 扫描件 (纯图片 PDF) 会抽不到文本, 此时**明确报错**而非静默产出空文档。
  - 体积与字符数有上限, 防止超大文件打爆内存与向量库。

🏭 Java 对标: 文件解析适配层 (POI / PDFBox 封装), 按扩展名路由 + Context
"""

from __future__ import annotations

import codecs
import csv
import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Protocol

from app.config import Settings
from app.core.exceptions import DocumentParseError

logger = logging.getLogger(__name__)

__all__ = [
    "DocumentParseError",
    "DocumentLoader",
    "DocumentParser",
    "ExtractedDocument",
    "PlainTextParser",
    "PdfParser",
    "DocxParser",
    "SpreadsheetParser",
    "CsvParser",
    "create_document_loader",
    "decode_text_bytes",
    "default_parsers",
    "normalize_text",
    "render_rows_as_text",
]


@dataclass
class ExtractedDocument:
    """解析结果 — 纯文本 + 溯源元数据"""

    text: str
    metadata: dict = field(default_factory=dict)


# ── 文本清洗 ─────────────────────────────────────────────
# 控制字符 (保留 \t \n): pypdf 对某些字体/编码会产出 \x00 等垃圾
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
# PDF 排版常把字距误判成空格, 于是汉字被拆开 ("理 赔 流 程"); 中文场景应合并
_HAN_GAP = re.compile(r"(?<=[\u4e00-\u9fff])[ \t]+(?=[\u4e00-\u9fff])")
_BLANK_LINES = re.compile(r"\n{3,}")


def normalize_text(text: str) -> str:
    """清洗解析出的文本 — 真实 PDF / Word 抽出来的文本几乎都需要过一遍

    - 统一换行 (CRLF / CR → LF)
    - 去掉 NUL 与控制字符
    - 不换行空格 / 全角空格 → 普通空格
    - 合并被误判成空格的汉字间隔 (中文 PDF 最常见的脏数据)
    - 压缩连续空白与空行, 逐行去首尾空白
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL_CHARS.sub("", text)
    text = text.replace("\u00a0", " ").replace("\u3000", " ")
    text = _HAN_GAP.sub("", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = _BLANK_LINES.sub("\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


# ── 编码识别 ─────────────────────────────────────────────
# 顺序: 显式 BOM → utf-8 → gb18030 (GBK/GB2312 的超集, 中文资料兜底)
_TEXT_ENCODINGS = ("utf-8", "gb18030")


def decode_text_bytes(data: bytes) -> tuple[str, str]:
    """中文友好的字节 → 文本解码。

    返回 (文本, 实际使用的编码名)。显式 BOM 优先; 无 BOM 时依次尝试 utf-8 /
    gb18030; 全部失败则用 gb18030 容错解码 (中文场景下比 utf-8 replace 乱码更少)。
    """
    if data.startswith(codecs.BOM_UTF8):
        return data.decode("utf-8-sig"), "utf-8-sig"
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16"), "utf-16"
    for enc in _TEXT_ENCODINGS:
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return data.decode("gb18030", errors="replace"), "gb18030(replace)"


# ── 表格 → 文本 ──────────────────────────────────────────


def render_rows_as_text(rows: list[list[str]]) -> str:
    """表格行渲染成「表头: 值」的文本行。

    第一行视作表头; 全空行跳过。每行生成 "列A: 值1; 列B: 值2" —— 带列名上下文,
    切块后语义不丢, 也便于 BM25 命中列名关键词。
    """
    cleaned = [
        [("" if c is None else str(c).strip()) for c in row] for row in rows
    ]
    cleaned = [r for r in cleaned if any(r)]
    if not cleaned:
        return ""
    header, *body = cleaned
    if not body:  # 只有表头, 原样输出
        return " | ".join(header)
    lines: list[str] = []
    for row in body:
        pairs = []
        for i, cell in enumerate(row):
            if not cell:
                continue
            col = header[i] if i < len(header) and header[i] else ""
            pairs.append(f"{col}: {cell}" if col else cell)
        if pairs:
            lines.append("; ".join(pairs))
    return "\n".join(lines)


# ── 各格式解析器 ─────────────────────────────────────────


class DocumentParser(Protocol):
    """单一格式解析器协议"""

    extensions: tuple[str, ...]

    def parse(self, filename: str, data: bytes) -> ExtractedDocument:
        ...


class PlainTextParser:
    """纯文本 / Markdown (含编码自动识别)"""

    extensions = (".txt", ".md", ".markdown", ".log")

    def parse(self, filename: str, data: bytes) -> ExtractedDocument:
        text, encoding = decode_text_bytes(data)
        if not text.strip():
            raise DocumentParseError("文本文件内容为空")
        return ExtractedDocument(
            text=text,
            metadata={"ext": Path(filename).suffix.lower(), "encoding": encoding},
        )


class PdfParser:
    """PDF — 逐页抽取文本 (pypdf)

    扫描件 (纯图片 PDF) 抽不到文本, 明确报错而非产出空文档。
    """

    extensions = (".pdf",)

    def parse(self, filename: str, data: bytes) -> ExtractedDocument:
        try:
            from pypdf import PdfReader
        except ImportError as e:  # pragma: no cover - 依赖缺失分支
            raise DocumentParseError(
                "解析 PDF 需要 pypdf, 请执行: pip install pypdf"
            ) from e

        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                # 很多 PDF 只设了权限密码 (空用户密码), 先试空密码
                try:
                    reader.decrypt("")
                except Exception as e:
                    raise DocumentParseError("PDF 已加密, 无法解析") from e
            pages = [(page.extract_text() or "") for page in reader.pages]
        except DocumentParseError:
            raise
        except Exception as e:
            raise DocumentParseError(
                f"PDF 解析失败: {type(e).__name__}: {e}"
            ) from e

        with_text = [t.strip() for t in pages if t.strip()]
        if not with_text:
            raise DocumentParseError(
                "PDF 未抽到任何文本, 可能是扫描件 (图片型 PDF)。"
                "本解析器不含 OCR, 请先做 OCR 或改用文本版 PDF。"
            )
        return ExtractedDocument(
            text="\n\n".join(with_text),
            metadata={
                "ext": ".pdf",
                "pages": len(pages),
                "pages_with_text": len(with_text),
            },
        )


class DocxParser:
    """Word (.docx) — 段落 + 表格 (python-docx)

    仅支持 .docx; 老式 .doc 是二进制复合文档, 需先另存为 .docx。
    """

    extensions = (".docx",)

    def parse(self, filename: str, data: bytes) -> ExtractedDocument:
        try:
            import docx  # python-docx
        except ImportError as e:  # pragma: no cover - 依赖缺失分支
            raise DocumentParseError(
                "解析 Word 需要 python-docx, 请执行: pip install python-docx"
            ) from e

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as e:
            raise DocumentParseError(
                "Word 解析失败 (仅支持 .docx; 老式 .doc 请先另存为 .docx): "
                f"{type(e).__name__}: {e}"
            ) from e

        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        tables = list(document.tables)
        for table in tables:
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            rendered = render_rows_as_text(rows)
            if rendered:
                parts.append(rendered)
        text = "\n".join(parts)
        if not text.strip():
            raise DocumentParseError("Word 文档未抽到任何文本")
        return ExtractedDocument(
            text=text,
            metadata={
                "ext": ".docx",
                "paragraphs": len(document.paragraphs),
                "tables": len(tables),
            },
        )


class SpreadsheetParser:
    """Excel (.xlsx / .xlsm) — 逐工作表抽成行文本 (openpyxl)"""

    extensions = (".xlsx", ".xlsm")

    def parse(self, filename: str, data: bytes) -> ExtractedDocument:
        try:
            from openpyxl import load_workbook
        except ImportError as e:  # pragma: no cover - 依赖缺失分支
            raise DocumentParseError(
                "解析 Excel 需要 openpyxl, 请执行: pip install openpyxl"
            ) from e

        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise DocumentParseError(
                "Excel 解析失败 (仅支持 .xlsx/.xlsm; 老式 .xls 请另存为 .xlsx): "
                f"{type(e).__name__}: {e}"
            ) from e

        sheet_names = list(workbook.sheetnames)
        blocks: list[str] = []
        for worksheet in workbook.worksheets:
            rows = [
                ["" if v is None else str(v) for v in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            rendered = render_rows_as_text(rows)
            if rendered:
                blocks.append(f"【工作表: {worksheet.title}】\n{rendered}")
        workbook.close()

        text = "\n\n".join(blocks)
        if not text.strip():
            raise DocumentParseError("Excel 未抽到任何内容")
        return ExtractedDocument(
            text=text,
            metadata={"ext": Path(filename).suffix.lower(), "sheets": len(sheet_names)},
        )


class CsvParser:
    """CSV / TSV — 编码自适应 + 表头感知"""

    extensions = (".csv", ".tsv")

    def parse(self, filename: str, data: bytes) -> ExtractedDocument:
        text_in, encoding = decode_text_bytes(data)
        delimiter = "\t" if Path(filename).suffix.lower() == ".tsv" else ","
        rows = [
            [(cell or "").strip() for cell in row]
            for row in csv.reader(io.StringIO(text_in), delimiter=delimiter)
        ]
        rendered = render_rows_as_text(rows)
        if not rendered.strip():
            raise DocumentParseError("CSV/TSV 未抽到任何内容")
        return ExtractedDocument(
            text=rendered,
            metadata={
                "ext": Path(filename).suffix.lower(),
                "encoding": encoding,
                "rows": len(rows),
            },
        )


def default_parsers() -> list[DocumentParser]:
    """默认解析器集合 (按扩展名路由)"""
    return [
        PlainTextParser(),
        PdfParser(),
        DocxParser(),
        SpreadsheetParser(),
        CsvParser(),
    ]


class DocumentLoader:
    """文档解析入口 — 按扩展名路由 + 体积/长度兜底

    🏭 Java 对标: 一组 Strategy + 一个 Context (扩展名 → 解析器)
    """

    def __init__(
        self,
        parsers: list[DocumentParser] | None = None,
        max_bytes: int = 20 * 1024 * 1024,
        max_chars: int = 2_000_000,
    ):
        self._parsers = parsers if parsers is not None else default_parsers()
        self._by_ext: dict[str, DocumentParser] = {
            ext: parser for parser in self._parsers for ext in parser.extensions
        }
        self._max_bytes = max_bytes
        self._max_chars = max_chars

    def supported_extensions(self) -> list[str]:
        """当前支持的扩展名 (用于前端提示与错误信息)"""
        return sorted(self._by_ext)

    def load(self, filename: str, data: bytes) -> ExtractedDocument:
        """解析文件 → ExtractedDocument; 失败抛 DocumentParseError"""
        ext = Path(filename).suffix.lower()
        parser = self._by_ext.get(ext)
        if parser is None:
            raise DocumentParseError(
                f"不支持的文件类型 {ext or '(无扩展名)'}; "
                f"当前支持: {', '.join(self.supported_extensions())}"
            )
        if not data:
            raise DocumentParseError("文件内容为空")
        if len(data) > self._max_bytes:
            raise DocumentParseError(
                f"文件过大 ({len(data) / 1e6:.1f} MB, 上限 "
                f"{self._max_bytes / 1e6:.0f} MB)"
            )

        doc = parser.parse(filename, data)
        raw_chars = len(doc.text)
        doc.text = normalize_text(doc.text)
        if not doc.text:
            raise DocumentParseError(
                f"解析后无有效文本 (清洗后为空), 请检查 {ext} 文件内容"
            )

        truncated = False
        if len(doc.text) > self._max_chars:
            doc.text = doc.text[: self._max_chars]
            truncated = True
        doc.metadata.update(
            {
                "filename": filename,
                "bytes": len(data),
                "chars_raw": raw_chars,
                "chars": len(doc.text),
                "truncated": truncated,
            }
        )
        logger.info(
            "文档解析完成 filename=%s ext=%s chars=%d truncated=%s",
            filename, ext, len(doc.text), truncated,
        )
        return doc


def create_document_loader(settings: Settings) -> DocumentLoader:
    """文档解析工厂 — 上限走配置 (SOP_QA_DOC_MAX_BYTES / _CHARS)"""
    return DocumentLoader(
        max_bytes=settings.doc_max_bytes,
        max_chars=settings.doc_max_chars,
    )
